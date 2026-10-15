from io import BytesIO
from pathlib import Path
import streamlit as st
from docx import Document
from docx.shared import RGBColor

st.set_page_config(page_title="Generator SPPD KONI", page_icon="📄", layout="wide")

TEMPLATE = Path(__file__).parent / "template_sppd.docx"

st.title("Generator SPPD — KONI Kabupaten Tangerang")
st.caption("Isi kolom di bawah, lalu unduh dokumen Word dengan format berdasarkan templat asli.")

st.info(
    "Aplikasi ini menggunakan file Word asli sebagai templat. "
    "Tata letak tabel, garis, ukuran halaman, dan teks standar dipertahankan sejauh memungkinkan. "
    "Periksa hasil di Microsoft Word sebelum dicetak."
)

def replace_cell_text(cell, value, color="EE0000"):
    """
    Mengganti isi cell dengan nilai baru sambil mempertahankan properti paragraf
    dan format run pertama. Gunakan hanya untuk sel jawaban yang memang ingin diganti.
    """
    paragraphs = cell.paragraphs
    if not paragraphs:
        p = cell.add_paragraph()
    else:
        p = paragraphs[0]

    # Simpan format run pertama bila tersedia
    source_run = next((r for para in paragraphs for r in para.runs), None)
    font_name = source_run.font.name if source_run else "Arial"
    font_size = source_run.font.size if source_run else None
    bold = source_run.bold if source_run else None
    italic = source_run.italic if source_run else None

    # Bersihkan teks semua paragraf dalam sel, tanpa menghapus properti sel/tabel
    for para in paragraphs:
        for run in para.runs:
            run.text = ""
    run = p.runs[0] if p.runs else p.add_run()
    run.text = str(value or "")
    run.font.name = font_name
    if font_size:
        run.font.size = font_size
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)

def get_template():
    if not TEMPLATE.exists():
        st.error("File template_sppd.docx tidak ditemukan. Letakkan bersama app.py.")
        st.stop()
    return Document(str(TEMPLATE))

with st.form("sppd_form"):
    st.subheader("I. Perjalanan berangkat dan kembali")
    a, b = st.columns(2)
    with a:
        asal = st.text_input("Berangkat dari", "Tangerang")
        tujuan = st.text_input("Ke / keperluan", "Try Out Catur")
        tanggal_berangkat = st.text_input("Tanggal berangkat", "17 Oktober 2026")
        pejabat_berwenang = st.text_input("Pejabat yang berwenang", "Sekretaris KONI Kabupaten Tangerang")
        nama_pejabat = st.text_input("Nama pejabat", "Adang Akbarudin")
    with b:
        tiba_kembali = st.text_input("Tiba di (bagian I)", "Tangerang")
        tanggal_kembali = st.text_input("Tanggal tiba", "18 Oktober 2026")
        kantor_penerima = st.text_input("Yang menerima / kantor", "KONI Kabupaten Tangerang")
        nama_penerima = st.text_input("Nama penerima", "Adang Akbarudin")
        jabatan_penerima = st.text_input("Jabatan penerima", "Sekretaris")

    st.subheader("II. Lokasi perjalanan")
    c, d = st.columns(2)
    with c:
        lokasi_ii_berangkat = st.text_area("Lokasi tiba (kolom kiri bagian II)", "Piset Square. Jl. Pelajar Pejuang No. 119 Bandung")
        tanggal_ii_berangkat = st.text_input("Tanggal (kolom kiri bagian II)", "17 Oktober 2026")
        penerima_ii_kiri = st.text_input("Yang menerima (kiri II)", "")
        nama_ii_kiri = st.text_input("Nama (kiri II)", "")
        jabatan_ii_kiri = st.text_input("Jabatan (kiri II)", "")
    with d:
        lokasi_ii_tiba = st.text_area("Lokasi tiba (kolom kanan bagian II)", "Piset Square. Jl. Pelajar Pejuang No. 119 Bandung")
        tanggal_ii_tiba = st.text_input("Tanggal (kolom kanan bagian II)", "18 Oktober 2026")
        penerima_ii_kanan = st.text_input("Yang menerima (kanan II)", "")
        nama_ii_kanan = st.text_input("Nama (kanan II)", "")
        jabatan_ii_kanan = st.text_input("Jabatan (kanan II)", "")

    st.subheader("III. Catatan perjalanan tambahan")
    e, f = st.columns(2)
    with e:
        lokasi_iii_kiri = st.text_input("Lokasi (kiri III)", "")
        tanggal_iii_kiri = st.text_input("Tanggal (kiri III)", "")
        penerima_iii_kiri = st.text_input("Yang menerima (kiri III)", "")
        nama_iii_kiri = st.text_input("Nama (kiri III)", "")
        jabatan_iii_kiri = st.text_input("Jabatan (kiri III)", "")
    with f:
        lokasi_iii_kanan = st.text_input("Lokasi (kanan III)", "")
        tanggal_iii_kanan = st.text_input("Tanggal (kanan III)", "")
        penerima_iii_kanan = st.text_input("Yang menerima (kanan III)", "")
        nama_iii_kanan = st.text_input("Nama (kanan III)", "")
        jabatan_iii_kanan = st.text_input("Jabatan (kanan III)", "")

    st.subheader("IV. Catatan perjalanan tambahan")
    g, h = st.columns(2)
    with g:
        lokasi_iv_kiri = st.text_input("Lokasi (kiri IV)", "")
        tanggal_iv_kiri = st.text_input("Tanggal (kiri IV)", "")
        penerima_iv_kiri = st.text_input("Yang menerima (kiri IV)", "")
        nama_iv_kiri = st.text_input("Nama (kiri IV)", "")
        jabatan_iv_kiri = st.text_input("Jabatan (kiri IV)", "")
    with h:
        lokasi_iv_kanan = st.text_input("Lokasi (kanan IV)", "")
        tanggal_iv_kanan = st.text_input("Tanggal (kanan IV)", "")
        penerima_iv_kanan = st.text_input("Yang menerima (kanan IV)", "")
        nama_iv_kanan = st.text_input("Nama (kanan IV)", "")
        jabatan_iv_kanan = st.text_input("Jabatan (kanan IV)", "")

    buat = st.form_submit_button("Buat dokumen Word", type="primary")

if buat:
    doc = get_template()
    if not doc.tables:
        st.error("Templat tidak memiliki tabel yang dikenali.")
        st.stop()
    table = doc.tables[0]

    # Pemetaan kolom jawaban pada tabel asli: indeks baris, indeks kolom.
    # Teks label hitam dan tanda baca tetap tidak disentuh.
    replacements = {
        (0, 3): "\n".join([asal, tujuan, tanggal_berangkat, pejabat_berwenang, nama_pejabat, "", "", "", "…………………………."]),
        (0, 5): "\n".join([tiba_kembali, tanggal_kembali, kantor_penerima, nama_penerima, jabatan_penerima, "", "", "", "…………………………."]),
        (1, 3): "\n\n".join([lokasi_ii_berangkat, tanggal_ii_berangkat, penerima_ii_kiri, nama_ii_kiri, jabatan_ii_kiri, "", "…………………………."]),
        (1, 5): "\n".join([lokasi_ii_tiba, tanggal_ii_tiba, penerima_ii_kanan, nama_ii_kanan, jabatan_ii_kanan, "", "…………………………."]),
        (2, 3): "\n".join([lokasi_iii_kiri, tanggal_iii_kiri, penerima_iii_kiri, nama_iii_kiri, jabatan_iii_kiri, "", "", "…………………………."]),
        (2, 5): "\n".join([lokasi_iii_kanan, tanggal_iii_kanan, penerima_iii_kanan, nama_iii_kanan, jabatan_iii_kanan, "", "", "…………………………."]),
        (3, 3): "\n".join([lokasi_iv_kiri, tanggal_iv_kiri, penerima_iv_kiri, nama_iv_kiri, jabatan_iv_kiri, "", "", "…………………………."]),
        (3, 5): "\n".join([lokasi_iv_kanan, tanggal_iv_kanan, penerima_iv_kanan, nama_iv_kanan, jabatan_iv_kanan, "", "", "…………………………."]),
    }
    for (r, c), value in replacements.items():
        if r < len(table.rows) and c < len(table.rows[r].cells):
            replace_cell_text(table.rows[r].cells[c], value)

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    st.success("Dokumen berhasil dibuat. Silakan unduh dan periksa sebelum dicetak.")
    st.download_button(
        "⬇️ Unduh SPPD (.docx)",
        data=out.getvalue(),
        file_name="SPPD_KONI_Kabupaten_Tangerang.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

with st.expander("Cara menjalankan aplikasi"):
    st.markdown("""
1. Instal Python 3.10 atau lebih baru.
2. Buka terminal di folder aplikasi.
3. Jalankan `pip install -r requirements.txt`.
4. Jalankan `streamlit run app.py`.
5. Browser akan membuka halaman aplikasi.
""")
